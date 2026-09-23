"""
FaceAttend Presentation Generator
Generates a 15-slide widescreen modern PowerPoint (.pptx) presentation explaining
the FaceAttend system architecture, OpenCV computer vision pipeline, and full-stack workflows.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    COLOR_NAVY = RGBColor(15, 23, 42)        # #0F172A (Dark background / Text)
    COLOR_WHITE = RGBColor(255, 255, 255)    # #FFFFFF
    COLOR_PRIMARY = RGBColor(37, 99, 235)    # #2563EB (Brand Blue)
    COLOR_TEAL = RGBColor(13, 148, 136)      # #0D9488 (Student accent)
    COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B (Slate grey)
    COLOR_CARD_BG = RGBColor(248, 250, 252)  # #F8FAFC (Light card fill)
    COLOR_BORDER = RGBColor(226, 232, 240)   # #E2E8F0 (Card border)
    COLOR_GREEN = RGBColor(16, 185, 129)     # #10B981 (Success / Present)
    COLOR_ORANGE = RGBColor(245, 158, 11)    # #F59E0B (Warning / Late)

    def add_header(slide, title_text, category="FACEATTEND PROJECT PRESENTATION"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_PRIMARY

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_NAVY

    def add_footer(slide, current_slide, total_slides=15):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.4))
        tf = footer_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"FaceAttend — Automated Biometric Attendance System  |  Slide {current_slide} of {total_slides}"
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_MUTED

    def add_card(slide, left, top, width, height, title, points, header_color=COLOR_PRIMARY):
        # Card Background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        # Card Text Box
        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.25), width - Inches(0.5), height - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = header_color
        p_title.space_after = Pt(12)

        for pt in points:
            p = tf.add_paragraph()
            p.text = "• " + pt
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_NAVY
            p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Hero Theme)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_NAVY
    bg1.line.color.rgb = COLOR_NAVY

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "FACEATTEND"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Face Recognition Attendance Management System"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE
    p2.space_after = Pt(20)

    p3 = tf1.add_paragraph()
    p3.text = "A Complete Technical Walkthrough of Architecture, OpenCV Biometrics, and Dual-Portal Implementation"
    p3.font.size = Pt(15)
    p3.font.color.rgb = RGBColor(148, 163, 184)
    p3.space_after = Pt(36)

    p4 = tf1.add_paragraph()
    p4.text = "Technology: Python Flask  |  OpenCV (Haar Cascade & LBPH)  |  SQLite  |  Bootstrap 5 & Chart.js"
    p4.font.size = Pt(12)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_GREEN

    # -------------------------------------------------------------
    # SLIDE 2: Project Overview & Objectives
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Project Overview & Core Objectives")
    add_footer(s2, 2)

    add_card(s2, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.9), "The Challenge", [
        "Manual roll-call attendance is slow and error-prone.",
        "Paper registers or RFID cards allow buddy punching (proxy attendance).",
        "Cloud face-recognition APIs raise severe privacy concerns and require constant internet connectivity.",
        "Difficulty for students to verify personal attendance records in real time."
    ], COLOR_ORANGE)

    add_card(s2, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.9), "The Solution", [
        "FaceAttend: Automated, contactless facial biometric attendance system.",
        "Processes faces 100% locally and offline on Windows without third-party cloud leaks.",
        "Instant recognition in < 1 second using standard webcam hardware.",
        "Strict anti-proxy validation: only matching student biometrics are accepted."
    ], COLOR_PRIMARY)

    add_card(s2, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.9), "Core Goals", [
        "Zero Duplicates: Guaranteed 1 attendance record per student per day.",
        "Dual Separate Dashboards: Dedicated interfaces for Admin and Student.",
        "Self-Attendance: Allows students to verify and mark attendance from their profile.",
        "Audit & Analytics: Complete reporting with CSV, Excel, and charts."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 3: Technology Stack & Ecosystem
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Full-Stack Technology Architecture")
    add_footer(s3, 3)

    add_card(s3, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.9), "Backend & Server", [
        "Python 3.10+ / 3.14: Core programming language.",
        "Flask 3.1: Lightweight, modular WSGI web application framework.",
        "Blueprint Architecture: Isolated routing namespaces (/admin, /user, /api, /auth).",
        "Werkzeug Security: Robust PBKDF2 with SHA-256 password hashing."
    ], COLOR_PRIMARY)

    add_card(s3, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.9), "Computer Vision Engine", [
        "OpenCV (opencv-contrib-python): Local computer vision library.",
        "Haar Cascade Classifier: High-speed real-time face detection.",
        "LBPH Face Recognizer: Texture-based histogram classification.",
        "NumPy: High-performance array and frame transformation."
    ], COLOR_TEAL)

    add_card(s3, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.9), "Database & Frontend", [
        "SQLite 3: Embedded ACID relational database with foreign keys & indexes.",
        "HTML5 & WebRTC: Native navigator.mediaDevices for browser camera streaming.",
        "Bootstrap 5.3 & CSS3: Fully responsive layouts with dark/light themes.",
        "Chart.js 4.4 & OpenPyXL: Dynamic visualization and styled Excel exports."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 4: End-to-End Workflow Diagram
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "End-to-End System Workflow")
    add_footer(s4, 4)

    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "Face Registration Flow (Admin)", [
        "1. Admin selects student from enrolled registry.",
        "2. Browser activates webcam via HTML5 getUserMedia.",
        "3. Real-time Haar Cascade checks for single face.",
        "4. Captures 5 to 6 normalized face crops (200x200 px).",
        "5. Crops saved to database/faces/<student_id>/.",
        "6. LBPH model trains automatically and writes face_model.xml.",
        "7. Student biometric status changes to 'Enrolled'."
    ], COLOR_PRIMARY)

    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Attendance Marking Flow (Live)", [
        "1. Camera streams frames to /api/face/recognize.",
        "2. Haar Cascade detects face and extracts normalized ROI.",
        "3. LBPH model predicts student ID and confidence distance.",
        "4. Validation: If distance > threshold -> Reject as 'Unknown Face'.",
        "5. Anti-Proxy: If student portal, verifies scanned face matches logged-in user.",
        "6. Duplicate Check: Checks if student already marked today.",
        "7. Time Cutoff: Compares with 09:30 AM -> Flags 'Present' or 'Late'.",
        "8. Records saved to SQLite with audio-visual feedback."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 5: Technique 1 — Face Detection (Haar Cascade)
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Technique 1: OpenCV Haar Cascade Face Detection")
    add_footer(s5, 5)

    add_card(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "How Haar Cascade Works", [
        "Uses machine-learning object detection proposed by Paul Viola and Michael Jones.",
        "Haar-like Features: Rectangular digital image features that analyze contrast gradients (e.g. eyes are darker than forehead).",
        "Integral Images: Enables rapid calculation of rectangular feature sums in constant time O(1).",
        "Adaboost Classifier: Selects the most critical facial features from thousands of candidates.",
        "Cascade of Stages: Quickly eliminates non-face background windows in early stages to run in real-time."
    ], COLOR_PRIMARY)

    add_card(s5, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "In-App Implementation & Safety Checks", [
        "Model File: haarcascade_frontalface_default.xml loaded locally.",
        "Parameters: scaleFactor=1.2, minNeighbors=5, minSize=(80, 80).",
        "Three-Way Validation Gate:",
        "  • 0 Faces: 'No face detected. Please position face.'",
        "  • >1 Faces: 'Multiple faces detected. Ensure only one person is visible.'",
        "  • 1 Face: Bounding box extracted with safety margin."
    ], COLOR_TEAL)

    # -------------------------------------------------------------
    # SLIDE 6: Technique 2 — Image Preprocessing & Normalization
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Technique 2: Face Normalization & Lighting Invariance")
    add_footer(s6, 6)

    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "The Illumination Challenge", [
        "Raw webcam images suffer from environmental variations:",
        "  • Harsh directional sunlight vs dim classroom lighting.",
        "  • Variable background noise and color shifts.",
        "  • Differing distances from camera producing different face sizes.",
        "Without normalization, recognition algorithms fail due to shadow intensity changes rather than facial differences."
    ], COLOR_ORANGE)

    add_card(s6, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Normalization Pipeline in FaceAttend", [
        "1. Grayscale Conversion: cv2.cvtColor(img, COLOR_BGR2GRAY) removes redundant color channels and focuses on facial geometry.",
        "2. Histogram Equalization: cv2.equalizeHist(gray) stretches pixel contrast and flattens uneven lighting/shadows.",
        "3. ROI Cropping with Margin: Adds 5% border around face bounding box to preserve jawline and hairline.",
        "4. Standard Resizing: Rescales all face crops to exact 200x200 matrix."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 7: Technique 3 — Face Recognition (LBPH Algorithm)
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Technique 3: LBPH Face Recognition Algorithm")
    add_footer(s7, 7)

    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "What is LBPH?", [
        "LBPH = Local Binary Patterns Histograms.",
        "Invented for texture analysis, later adapted for highly robust face recognition (Ahonen et al., 2006).",
        "Unlike holistic methods (Eigenfaces) that look at the entire face at once, LBPH analyzes local micro-textures.",
        "Configured in FaceAttend: radius=1, neighbors=8, grid_x=8, grid_y=8."
    ], COLOR_PRIMARY)

    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "The 4-Step Mathematical Process", [
        "1. LBP Operator: Compares each pixel against its 8 neighbors. Neighbor >= Center ? 1 : 0. Produces an 8-bit binary number.",
        "2. Decimal Value: Converts binary number to decimal (0 to 255) for the pixel.",
        "3. Grid Division: Divides 200x200 face into an 8x8 grid (64 sub-regions).",
        "4. Histogram Concatenation: Calculates local histogram for each cell and concatenates them into one combined biometric feature vector."
    ], COLOR_TEAL)

    # -------------------------------------------------------------
    # SLIDE 8: Technique 4 — Chi-Square Distance & Thresholding
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Technique 4: Distance Matching & Unknown Face Rejection")
    add_footer(s8, 8)

    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "Chi-Square Distance Metric", [
        "When an unknown face is captured, its LBPH histogram is computed.",
        "OpenCV computes the Chi-Square distance between the query histogram and trained student prototypes.",
        "Distance Meaning in LBPH:",
        "  • Distance = 0: Exact theoretical match.",
        "  • Distance 40 - 70: Strong biometric match.",
        "  • Distance > 75: Mismatch or unregistered person."
    ], COLOR_PRIMARY)

    add_card(s8, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Threshold Decision Logic in FaceAttend", [
        "threshold = 75.0 (default, adjustable via Admin Settings slider).",
        "If distance <= threshold AND label in database:",
        "  -> RECOGNIZED: Identity confirmed, marks attendance.",
        "If distance > threshold OR unregistered face:",
        "  -> UNKNOWN FACE: Attendance is NOT recorded. Prompt to enroll.",
        "Confidence calculation: match_pct = max(0, 100 - distance)."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 9: Dual Dashboard Architecture & RBAC
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Dual Dashboard Architecture & Role-Based Access")
    add_footer(s9, 9)

    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "Strict Role-Based Separation", [
        "The system enforces two completely distinct dashboard layouts:",
        "1. Admin Dashboard (/admin/...): Institutional overview, student management, camera scanners, user roles, reports.",
        "2. Student Dashboard (/user/...): Student-friendly theme, personal metrics, self-attendance, history, certificate report.",
        "No Shared Single Dashboards: Navigation, templates, and backend controls are strictly isolated."
    ], COLOR_PRIMARY)

    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Backend Security Enforcement", [
        "Protected by Flask custom decorators in utils/auth_decorators.py:",
        "  • @admin_required: Verifies session['role'] == 'ADMIN'. If student tries /admin/..., rejects with 'Access Denied' and redirects to student home.",
        "  • @user_required: Verifies session['role'] == 'USER'.",
        "  • @login_required: Redirects unauthenticated requests to login.",
        "Permissions enforced at the controller level, not just by hiding UI buttons."
    ], COLOR_TEAL)

    # -------------------------------------------------------------
    # SLIDE 10: Admin Management Features
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Administrator Capabilities & Portal Tools")
    add_footer(s10, 10)

    add_card(s10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "Student & Biometric Management", [
        "Student Directory: Add, edit, delete, and view enrolled students with course, semester, department filters.",
        "Automatic User Provisioning: Creating a student automatically creates their matching portal login account.",
        "Register Face Page (/admin/register-face): Live webcam capture with sample progress counter (0 to 6) and auto-training.",
        "Delete Face Data: Admins can purge biometric samples and trigger model rebuild instantly."
    ], COLOR_PRIMARY)

    add_card(s10, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Master Scanner & User Accounts", [
        "Take Attendance Scanner (/admin/take-attendance): Continuous live kiosk camera recognizing students and displaying photo card.",
        "Attendance Master Records: Search, filter by date/dept/status, manual override, and audit deletion.",
        "User Management (/admin/users): Create users, disable/activate accounts, reset passwords.",
        "Global Settings: Configure institute title, late cutoff time (e.g. 09:30 AM), and recognition sensitivity."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 11: Student Self-Attendance Feature
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Student Self-Attendance & Anti-Proxy Verification")
    add_footer(s11, 11)

    add_card(s11, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "How Self-Attendance Works", [
        "Accessible via /user/mark-attendance or from Student Profile / Dashboard.",
        "Student activates laptop or tablet webcam in their personal browser.",
        "Scanner streams frame to /api/face/self-attendance.",
        "System recognizes face AND performs identity cross-check:",
        "  Is recognized_student_id == session['student_id']?"
    ], COLOR_TEAL)

    add_card(s11, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Anti-Proxy & Security Measures", [
        "1. Identity Match: If another student's face is scanned, system alerts: 'Face Mismatch! Scanned face does not match your profile.'",
        "2. Unknown Person: If an unregistered person scans, attendance is denied.",
        "3. Real-Time Status Update: Instant UI celebration card, sound chime, and today's status badge turns green.",
        "4. Cutoff Check: Compares current time with cutoff to log Present or Late."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 12: Duplicate Prevention Architecture
    # -------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Duplicate Attendance Prevention Architecture")
    add_footer(s12, 12)

    add_card(s12, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "The Problem: Double Dipping", [
        "In camera-based systems, a student often remains in front of the lens for multiple seconds or returns later in the day.",
        "Without protection, hundreds of duplicate rows flood the database, distorting attendance rates and reports.",
        "Challenge: Prevent duplicates gracefully without crashing the camera stream."
    ], COLOR_ORANGE)

    add_card(s12, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Two-Tier Defense in FaceAttend", [
        "Tier 1: Application Logic Check",
        "  SELECT id FROM attendance WHERE student_id = ? AND attendance_date = ?",
        "  If found: Returns 'already_marked' status -> Shows 'Attendance Already Marked Today' alert.",
        "Tier 2: Database Hard Constraint",
        "  CONSTRAINT unique_daily_attendance UNIQUE (student_id, attendance_date)",
        "  Even in race conditions, SQLite rejects duplicate daily insertions at the database engine level."
    ], COLOR_PRIMARY)

    # -------------------------------------------------------------
    # SLIDE 13: Analytics, Reporting & Exports
    # -------------------------------------------------------------
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Analytics, Reports & Export Formats")
    add_footer(s13, 13)

    add_card(s13, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "Dynamic Real-Time Charts", [
        "All chart metrics calculated dynamically from SQLite (never hardcoded!):",
        "  • 7-Day Attendance Trend: Multi-line Chart.js showing Present, Late, and Absent daily curves.",
        "  • Department Distribution: Bar chart comparing attendance volume across engineering branches.",
        "  • Personal Attendance Donut: Student's personal attendance ratio with 75% compliance threshold."
    ], COLOR_PRIMARY)

    add_card(s13, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Multi-Channel Export Options", [
        "1. CSV Export: Instant lightweight download of raw attendance logs for custom spreadsheets.",
        "2. Excel Export (.xlsx): Styled workbooks created using openpyxl with formatted headers and color-coded status badges.",
        "3. Print-Ready Transcripts: Clean window.print() CSS styles for official student attendance certificates."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 14: Data Privacy & Security Guardrails
    # -------------------------------------------------------------
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Security Architecture & Biometric Privacy")
    add_footer(s14, 14)

    add_card(s14, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "Biometric Privacy Protections", [
        "100% On-Premises: Face crops, LBPH model files, and encodings stay entirely on the local server.",
        "Zero Cloud Leaks: No student photos or biometric tokens sent to third-party commercial APIs.",
        "Right to Erasure (GDPR / Privacy Compliance): Admins can purge biometric records on request.",
        "Encodings Protected: Raw biometric vectors are never exposed to frontend JavaScript."
    ], COLOR_TEAL)

    add_card(s14, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Application Security Defenses", [
        "Password Hashing: Stored using PBKDF2 with SHA-256 (Werkzeug security).",
        "SQL Injection Prevention: 100% of SQLite database queries use parameterized SQL statements (?, ?).",
        "HttpOnly Sessions: Prevents session hijacking via cross-site scripting.",
        "Automated Test Coverage: 14/14 automated unit tests verifying authentication, biometrics, and RBAC."
    ], COLOR_PRIMARY)

    # -------------------------------------------------------------
    # SLIDE 15: Conclusion & Future Scope
    # -------------------------------------------------------------
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, "Summary & Future Roadmap")
    add_footer(s15, 15)

    add_card(s15, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), "What We Achieved", [
        "Complete full-stack working biometric attendance system.",
        "High accuracy, rapid recognition (< 1s) with zero external GPU requirements.",
        "Dual role-based dashboards with total separation between Admin and Student.",
        "Complete duplicate prevention and anti-proxy self-attendance verification.",
        "Verified with 14 automated unit tests."
    ], COLOR_GREEN)

    add_card(s15, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Future Roadmap", [
        "Liveness Detection: Implement blink detection or infrared depth sensing to prevent photo spoofing.",
        "RTSP IP Camera Integration: Multi-camera continuous hallway scanning.",
        "Mobile App: Progressive Web App (PWA) with geofencing check-in.",
        "Automated Notifications: Email / SMS alerts to students with low attendance (< 75%)."
    ], COLOR_PRIMARY)

    # Save presentation
    output_path = os.path.join(os.path.dirname(__file__), "FaceAttend_Project_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation created successfully at: {output_path}")

if __name__ == '__main__':
    create_presentation()
